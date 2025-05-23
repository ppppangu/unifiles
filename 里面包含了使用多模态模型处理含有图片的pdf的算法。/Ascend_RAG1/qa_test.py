from docx import Document
from query_data import query_rag_test
import pandas as pd
import os


# 提取出docx文件(不包含答案的docx文件)中的测试问题
def extract_docx_questions_with_font_size(file_path, font_size=14):
    # 打开 .docx 文件
    doc = Document(file_path)

    # 存储提取的问题
    questions = []

    # 遍历文档中的段落
    for para in doc.paragraphs:
        # 检查段落中的每个 run（run 是具有相同样式的文本片段）
        for run in para.runs:
            # 检查字体大小是否为指定的值
            if run.font.size and run.font.size.pt == font_size:
                if para.text.endswith("？") or para.text.endswith("?"):
                    questions.append(para.text)
                    break  # 如果找到符合条件的 run，跳过该段落的其余部分

    return questions


# 提取出表格中的问题
def extract_xlsx_questions(file_path):

    questions = []
    # 使用 pandas 读取 Excel 文件
    df = pd.read_excel(file_path)

    # 按行遍历数据
    for index, row in df.iterrows():
        if index == 0 or index == 1:
            continue
        if len(row.values) == 3:
            question = row.values[1]
            questions.append(question)

    return questions


def write_docx(file_path, content):
    # 创建一个新的 Document 对象
    doc = Document()

    # 写入内容
    for text in content:
        doc.add_paragraph(text)  # 添加段落

    # 保存文档
    doc.save(file_path)
    print(f"文件已保存到: {file_path}")


# 对问题进行问答测试，将问题+解答存储到docx文件
if __name__ == "__main__":
    # 测试文件路径
    file_path_input = "data/test/question"
    # 测试结果输出路径
    file_path_output = "data/test/answer/问答测试结果.docx"
    answers = []
    questions = []

    # 提取字体大小为 14 的问题
    font_size = 14

    for file_name in os.listdir(file_path_input):
        file_path = os.path.join(file_path_input, file_name)

        if "xlsx" in file_path:
            questions += extract_xlsx_questions(file_path)
        elif "docx" in file_path:
            questions += extract_docx_questions_with_font_size(file_path, font_size)

    # 打印提取的问题
    for i, question in enumerate(questions, start=1):
        answer = query_rag_test(question)
        answer = f"{i}. " + question + "\n" + "回答： " + "\n" + answer + "\n" + "-------------------------------------"
        answers.append(answer)
        print(f"{i}. {question}")

    # 创建并写入 .docx 文件
    write_docx(file_path_output, answers)
